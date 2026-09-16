
def clean_en_date(txt):
    if not txt:
        return ""
    s = str(txt).strip().replace("،", " ").replace("م", "PM").replace("ص", "AM")
    import re
    s = re.sub(r'\s+', ' ', s)
    return s

from django.views.decorators.clickjacking import xframe_options_sameorigin
from django.shortcuts import get_object_or_404, render
from django.db import connection
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
import json

from .models import Violation


# أسماء الحقول التي يمكن تعديلها، مع دعم أسماء الواجهة العربية أيضًا.
FIELD_MAP = {
    "رقم المحضر": "violation_number",
    "المنسوب له المخالفة": "attributed_person",
    "موقع المخالفة": "violation_location",
    "موقع المخالفة (المنطقة، المدينة)": "violation_location",
    "اسم النظام أو اللائحة المستند عليها": "regulation",
    "اسم النظام أو اللائحة المستند عليها (وفق محضر ضبط المخالفة)": "regulation",
    "تكرار الإحالة": "referral_repetition",
    "محرر المحضر": "report_editor",
    "الفرع": "branch",
    "تاريخ تحرير المحضر": "report_date",
    "تاريخ تحرير المحضر (MM/DD/YYYY)": "report_date",
    "نوع المخالفة": "violation_type",
    "تاريخ إشعار المنسوب له المخالفة": "notified_date",
    "تاريخ رد المنسوب له المخالفة": "person_response_date",
    "تاريخ إحالتها إلى القطاع المختص": "sector_referral_date",
    "تاريخ رد القطاع المختص": "sector_response_date",
    "تاريخ إحالتها الى الأمانة": "secretariat_referral_date",
    "تاريخ رد الأمانة": "secretariat_response_date",
    "حالة المخالفة": "violation_status",
    "حالة المخالفة (MM/DD/YYYY)": "violation_status",
    "حالة المعاملة": "transaction_status",
    "حالة المعاملة (MM/DD/YYYY)": "transaction_status",
    "تاريخ إصدار قرار المخالفة من اللجنة": "committee_decision_date",
    "تاريخ إصدار قرار المخالفة من اللجنة (MM/DD/YYYY)": "committee_decision_date",
    "ملاحظات": "notes",
    "خط العرض": "latitude",
    "خط الطول": "longitude",
    "منشئ المخالفة": "created_by",
    "آخر تعديل بواسطة": "updated_by",
    "بيانات المحضر": "report_data",
    "المحضر": "report_data",
}

MODEL_FIELDS = {f.name for f in Violation._meta.fields}
CREATE_FIELDS = MODEL_FIELDS - {"id", "created_at", "updated_at"}
UPDATE_FIELDS = CREATE_FIELDS - {"created_by"}


def home(request):
    return render(request, "main_dashboard/alex.html", {"violations": Violation.objects.all()})


def add_violation(request):
    return render(request, "main_dashboard/alex.html")


def edit_violation(request, pk):
    return render(request, "main_dashboard/alex.html")


def geocode_search(request):
    return JsonResponse({})


def serialize_violation(v):
    return {
        f.name: getattr(v, f.name)
        for f in Violation._meta.fields
    }


def translate_payload(data, allowed_fields):
    translated = {}
    unknown = []
    for key, value in data.items():
        if key == "id":
            continue
        field = FIELD_MAP.get(key, key)
        if field not in allowed_fields:
            unknown.append(key)
            continue
        translated[field] = value
    if unknown:
        raise ValueError("حقول غير مسموح بها: " + ", ".join(map(str, unknown)))
    return translated


def api_violations(request):
    # SERA_DB_CONNECTION_DIAGNOSTIC_V1
    import logging
    _sera_logger = logging.getLogger("django.request")
    _sera_logger.warning(
        "SERA_DB_DIAGNOSTIC vendor=%s engine=%s host=%s",
        connection.vendor,
        connection.settings_dict.get("ENGINE"),
        connection.settings_dict.get("HOST") or "(local)",
    )
    if request.method != "GET":
        return JsonResponse({"ok": False, "error": "Method not allowed"}, status=405)
    results = list(Violation.objects.order_by("id").values())
    return JsonResponse({"ok": True, "count": len(results), "results": results})


@csrf_exempt
def api_violation_create(request):
    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "Method not allowed"}, status=405)
    try:
        data = json.loads(request.body or "{}")
        payload = translate_payload(data, CREATE_FIELDS)
        violation = Violation.objects.create(**payload)
        return JsonResponse({
            "ok": True,
            "result": serialize_violation(violation),
        }, status=201)
    except (json.JSONDecodeError, TypeError, ValueError) as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=400)
    except Exception as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=500)


@csrf_exempt
def api_violation_update(request, pk=None):
    if request.method not in {"PATCH", "PUT", "POST"}:
        return JsonResponse({"ok": False, "error": "Method not allowed"}, status=405)
    try:
        data = json.loads(request.body or "{}")
        v_id = pk or data.get("id")
        if not v_id:
            return JsonResponse({"ok": False, "error": "Missing violation id"}, status=400)

        violation = get_object_or_404(Violation, pk=v_id)
        payload = translate_payload(data, UPDATE_FIELDS)

        # آخر تعديل يُحفظ من الخادم حتى لا يعتمد على وقت الجهاز.
        payload["updated_at"] = timezone.now()
        if not payload.get("updated_by"):
            payload["updated_by"] = ""

        for field, value in payload.items():
            setattr(violation, field, value)
        violation.save()

        return JsonResponse({
            "ok": True,
            "result": serialize_violation(violation),
        })
    except (json.JSONDecodeError, TypeError, ValueError) as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=400)
    except Exception as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=500)


@csrf_exempt
def api_violation_delete(request, pk):
    if request.method not in {"POST", "DELETE"}:
        return JsonResponse({"ok": False, "error": "طريقة الطلب غير مسموح بها"}, status=405)
    try:
        violation = get_object_or_404(Violation, pk=pk)
        v_num = violation.violation_number or str(violation.pk)
        violation.delete()
        return JsonResponse({
            "ok": True,
            "message": f"تم حذف المخالفة رقم ({v_num}) بنجاح من قاعدة البيانات."
        })
    except Exception as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=500)


@csrf_exempt
def api_export_excel(request):
    import os
    import openpyxl
    from django.http import HttpResponse, JsonResponse
    from openpyxl.styles import PatternFill, Font
    from io import BytesIO
    import json

    if request.method != "POST":
        return JsonResponse({"ok": False, "error": "POST required"}, status=405)

    try:
        data = json.loads(request.body or "{}")
        rows = data.get("rows", [])

        template_path = os.path.join(os.path.dirname(__file__), "template_report.xlsx")
        wb = openpyxl.load_workbook(template_path)
        ws = wb['متابعة المحاضر']

        # مسح البيانات السابقة مع الاحتفاظ بالتنسيق
        for r in range(2, ws.max_row + 1):
            for c in range(1, 20):
                ws.cell(row=r, column=c).value = None

        col_keys = [
            'رقم المحضر',
            'المنسوب له المخالفة',
            'موقع المخالفة (المنطقة، المدينة)',
            'اسم النظام أو اللائحة المستند عليها (وفق محضر ضبط المخالفة)',
            'تكرار الإحالة',
            'محرر المحضر',
            'الفرع',
            'تاريخ تحرير المحضر (MM/DD/YYYY)',
            'نوع المخالفة',
            'تاريخ إشعار المنسوب له المخالفة (MM/DD/YYYY)',
            'تاريخ رد المنسوب له المخالفة (MM/DD/YYYY)',
            'تاريخ إحالتها إلى القطاع المختص  (MM/DD/YYYY)',
            'تاريخ رد القطاع المختص (MM/DD/YYYY)',
            'تاريخ إحالتها الى الأمانة (MM/DD/YYYY)',
            'تاريخ رد الأمانة (MM/DD/YYYY)',
            'حالة المخالفة (MM/DD/YYYY)',
            'حالة المعاملة (MM/DD/YYYY)',
            'تاريخ إصدار قرار المخالفة من اللجنة (MM/DD/YYYY)',
            'ملاحظات'
        ]

        pink_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        red_font = Font(name='SERA', size=16, color='9C0006', bold=True)

        for idx, item in enumerate(rows):
            r_num = 2 + idx
            ws.row_dimensions[r_num].height = 42
            for c_idx, key in enumerate(col_keys, start=1):
                val = item.get(key, '')
                cell = ws.cell(row=r_num, column=c_idx)
                cell.value = val

                # تمييز تكرار الإحالة عند وجود تكرار >= 2
                if c_idx == 5:
                    try:
                        rep = int(str(val).strip())
                        if rep >= 2:
                            cell.fill = pink_fill
                            cell.font = red_font
                    except (ValueError, TypeError):
                        pass

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = "attachment; filename*=UTF-8''SERA_متابعة_المحاضر.xlsx"
        return response

    except Exception as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=500)


@csrf_exempt
def api_sync_check(request):
    from django.db.models import Max
    from django.http import JsonResponse
    try:
        count = Violation.objects.count()
        latest_update = Violation.objects.aggregate(Max('updated_at'))['updated_at__max']
        latest_id = Violation.objects.aggregate(Max('id'))['id__max'] or 0
        up_str = latest_update.isoformat() if latest_update else ""
        
        return JsonResponse({
            "ok": True,
            "count": count,
            "latest_update": up_str,
            "latest_id": latest_id
        })
    except Exception as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=500)


# ==========================================================
# دالة توليد محضر SERA الرسمي فائق النقاء (Backend PDF Engine)
# ==========================================================
def _ar(text):
    if not text:
        return ""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        reshaped = arabic_reshaper.reshape(str(text))
        return get_display(reshaped)
    except Exception:
        return str(text)

from django.views.decorators.clickjacking import xframe_options_sameorigin

@xframe_options_sameorigin
def export_violation_official_pdf(request, pk):
    import os, io, pymupdf as fitz, arabic_reshaper
    from bidi.algorithm import get_display
    from django.http import HttpResponse, Http404
    from django.conf import settings
    from .models import Violation

    try:
        violation = Violation.objects.get(pk=pk)
    except Violation.DoesNotExist:
        raise Http404("المخالفة غير موجودة")

    p1_img = os.path.join(settings.BASE_DIR, 'main_dashboard', 'reports', 'page_1.png')
    p2_img = os.path.join(settings.BASE_DIR, 'main_dashboard', 'reports', 'page_2.png')

    font_paths = [
        os.path.join(settings.BASE_DIR, 'main_dashboard', 'static', 'main_dashboard', 'fonts', 'Amiri-Regular.ttf'),
        os.path.join(settings.BASE_DIR, 'assets', 'fonts', 'Amiri-Regular.ttf'),
        os.path.join(settings.BASE_DIR, 'main_dashboard', 'static', 'main_dashboard', 'fonts', 'Cairo-Regular.ttf')
    ]
    font_path = next((f for f in font_paths if os.path.exists(f)), None)

    doc = fitz.open()
    A4_RECT = fitz.Rect(0, 0, 595.28, 841.89)

    def _ar(txt):
        if not txt:
            return ""
        try:
            return get_display(arabic_reshaper.reshape(str(txt)))
        except Exception:
            return str(txt)

    def write_box(page, rect, text, fontsize=9.5, align=fitz.TEXT_ALIGN_CENTER, color=(0,0,0)):
        if not text:
            return
        page.insert_textbox(rect, _ar(text), fontsize=fontsize, fontname="amiri", fontfile=font_path, color=color, align=align)

    # Page 1
    page1 = doc.new_page(width=A4_RECT.width, height=A4_RECT.height)
    if os.path.exists(p1_img):
        page1.insert_image(A4_RECT, filename=p1_img)
    if font_path:
        page1.insert_font(fontname="amiri", fontfile=font_path)

    if getattr(violation, 'violation_number', None):
        write_box(page1, fitz.Rect(210, 202, 385, 218), f"({violation.violation_number})", fontsize=10)

    write_box(page1, fitz.Rect(248, 222, 384, 256), getattr(violation, 'attributed_person', '') or "", fontsize=9)
    write_box(page1, fitz.Rect(46, 222, 140, 256), str(getattr(violation, 'unified_number', '') or ""), fontsize=9.5)
    write_box(page1, fitz.Rect(248, 264, 384, 300), getattr(violation, 'violation_type', '') or "", fontsize=8.5)
    write_box(page1, fitz.Rect(46, 264, 140, 300), getattr(violation, 'violation_location', '') or "", fontsize=8.5)
    write_box(page1, fitz.Rect(248, 308, 384, 344), clean_en_date(getattr(violation, "report_date", "") or ""), fontsize=8.5)
    write_box(page1, fitz.Rect(46, 308, 140, 344), str(getattr(violation, 'incident_date', '') or ""), fontsize=8.5)

    regulation = getattr(violation, 'regulation', '') or getattr(violation, 'regulation_basis', '') or ""
    if regulation:
        write_box(page1, fitz.Rect(46, 395, 550, 560), regulation, fontsize=9, align=fitz.TEXT_ALIGN_RIGHT)

    facts = getattr(violation, 'facts', '') or getattr(violation, 'transaction_status', '') or ""
    if facts:
        write_box(page1, fitz.Rect(46, 715, 550, 810), facts, fontsize=9, align=fitz.TEXT_ALIGN_RIGHT)

    # Page 2
    page2 = doc.new_page(width=A4_RECT.width, height=A4_RECT.height)
    if os.path.exists(p2_img):
        page2.insert_image(A4_RECT, filename=p2_img)
    if font_path:
        page2.insert_font(fontname="amiri", fontfile=font_path)

    details = getattr(violation, 'details', '') or getattr(violation, 'description', '') or ""
    if details:
        write_box(page2, fitz.Rect(46, 215, 550, 335), details, fontsize=9, align=fitz.TEXT_ALIGN_RIGHT)

    damages = getattr(violation, 'damages', '') or getattr(violation, 'damage', '') or ""
    if damages:
        write_box(page2, fitz.Rect(46, 385, 550, 445), damages, fontsize=9, align=fitz.TEXT_ALIGN_RIGHT)

    documents = getattr(violation, 'documents', '') or ""
    if documents:
        write_box(page2, fitz.Rect(46, 485, 550, 555), documents, fontsize=9, align=fitz.TEXT_ALIGN_RIGHT)

    requests = getattr(violation, 'requests', '') or ""
    if requests:
        write_box(page2, fitz.Rect(46, 595, 550, 645), requests, fontsize=9, align=fitz.TEXT_ALIGN_RIGHT)

    defense = getattr(violation, 'defense', '') or ""
    if defense:
        write_box(page2, fitz.Rect(46, 685, 550, 755), defense, fontsize=9, align=fitz.TEXT_ALIGN_RIGHT)

    inspector = getattr(violation, 'inspector_name', '') or getattr(violation, 'report_editor', '') or ""
    if inspector:
        write_box(page2, fitz.Rect(328, 775, 472, 815), inspector, fontsize=10)

    pdf_bytes = doc.tobytes()
    doc.close()

    response = HttpResponse(pdf_bytes, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="sera_report_{violation.violation_number or pk}.pdf"'
    return response

